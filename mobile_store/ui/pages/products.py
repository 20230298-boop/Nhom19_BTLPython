"""
pages/products.py
─────────────────────────────────────────
Trang quản lý sản phẩm: CRUD, tìm kiếm, lọc, xuất Excel.
"""

import customtkinter as ctk
from tkinter import messagebox, filedialog
from datetime import datetime

from utils.constants import COLORS, FONT_FAMILY
from utils.helpers import format_currency, is_available
from ui.widgets import StyledButton, StyledTreeview

OPENPYXL = is_available("openpyxl")
if OPENPYXL:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment


class ProductsPage(ctk.CTkFrame):
    def __init__(self, master, db, user):
        super().__init__(master, fg_color=COLORS["bg"], corner_radius=0)
        self.db   = db
        self.user = user
        self._build()

    def _build(self):
        self._build_toolbar()
        self._build_table()
        self._load()

    def _build_toolbar(self):
        bar = ctk.CTkFrame(self, fg_color=COLORS["dark2"], corner_radius=0, height=60)
        bar.pack(fill="x")
        bar.pack_propagate(False)
        t = ctk.CTkFrame(bar, fg_color="transparent")
        t.pack(fill="both", expand=True, padx=20, pady=10)

        self.search_var = ctk.StringVar()
        ctk.CTkEntry(t, textvariable=self.search_var,
                     placeholder_text="🔍  Tìm kiếm sản phẩm...",
                     width=280, height=38,
                     fg_color=COLORS["card2"], border_color=COLORS["border"],
                     font=(FONT_FAMILY, 12),
                     text_color=COLORS["text"]).pack(side="left")
        self.search_var.trace("w", lambda *_: self._load())

        cats = ["Tất cả"] + [c["name"] for c in
                              self.db.fetchall("SELECT name FROM categories ORDER BY name")]
        self.cat_var = ctk.StringVar(value="Tất cả")
        ctk.CTkOptionMenu(t, variable=self.cat_var, values=cats,
                          command=lambda _: self._load(),
                          width=150, height=38,
                          fg_color=COLORS["card2"], button_color=COLORS["primary"],
                          font=(FONT_FAMILY, 11)).pack(side="left", padx=8)

        StyledButton(t, "Thêm mới", self._add,    style="success", icon="➕").pack(side="left", padx=4)
        StyledButton(t, "Sửa",      self._edit,   style="warning", icon="✏️").pack(side="left", padx=4)
        StyledButton(t, "Xóa",      self._delete, style="danger",  icon="🗑").pack(side="left", padx=4)
        StyledButton(t, "Làm mới",  self._load,   style="dark",    icon="🔄").pack(side="left", padx=4)
        if OPENPYXL:
            StyledButton(t, "Xuất Excel", self._export_excel, style="info", icon="📊").pack(side="right")

    def _build_table(self):
        cols   = ("code","name","category","brand","price_import","price_sell","qty","min_qty","warranty","status")
        heads  = ("Mã SP","Tên sản phẩm","Danh mục","Thương hiệu","Giá nhập","Giá bán","Tồn kho","Tối thiểu","BH(tháng)","Trạng thái")
        widths = [90, 220, 120, 100, 130, 130, 80, 80, 90, 100]
        self.tree = StyledTreeview(self, cols, heads, widths, height=25)
        self.tree.pack(fill="both", expand=True, padx=20, pady=(10, 5))

        self.status_var = ctk.StringVar()
        ctk.CTkLabel(self, textvariable=self.status_var,
                     font=(FONT_FAMILY, 11),
                     text_color=COLORS["text_muted"]).pack(anchor="w", padx=20, pady=(0, 8))

    def _load(self):
        self.tree.clear()
        kw  = self.search_var.get().strip()
        cat = self.cat_var.get()
        q = """SELECT p.*, c.name as cat_name FROM products p
               LEFT JOIN categories c ON p.category_id=c.id
               WHERE p.is_active=1"""
        params = []
        if kw:
            q += " AND (p.name LIKE ? OR p.code LIKE ? OR p.brand LIKE ?)"
            params += [f"%{kw}%"] * 3
        if cat != "Tất cả":
            q += " AND c.name=?"
            params.append(cat)
        q += " ORDER BY p.code"
        rows = self.db.fetchall(q, params)
        for r in rows:
            tags = ["low_stock"] if r["quantity"] <= r["min_quantity"] else []
            self.tree.insert_row((
                r["code"], r["name"], r["cat_name"] or "",
                r["brand"] or "",
                format_currency(r["import_price"]),
                format_currency(r["selling_price"]),
                r["quantity"], r["min_quantity"],
                r["warranty_months"],
                "✅ Đang bán",
            ), tags=tags)
        self.status_var.set(f"Tổng: {len(rows)} sản phẩm")

    def _add(self):
        from ui.dialogs.product_dialog import ProductDialog
        dlg = ProductDialog(self, self.db)
        self.wait_window(dlg)
        self._load()

    def _edit(self):
        sel = self.tree.get_selected()
        if not sel:
            messagebox.showwarning("Thông báo", "Chọn sản phẩm cần sửa!")
            return
        product = self.db.fetchone("SELECT * FROM products WHERE code=?", (sel[0],))
        if product:
            from ui.dialogs.product_dialog import ProductDialog
            dlg = ProductDialog(self, self.db, product)
            self.wait_window(dlg)
            self._load()

    def _delete(self):
        sel = self.tree.get_selected()
        if not sel:
            return
        if messagebox.askyesno("Xác nhận", f"Ẩn sản phẩm '{sel[1]}'?"):
            self.db.execute("UPDATE products SET is_active=0 WHERE code=?", (sel[0],))
            self.db.log_activity(self.user["id"], "DELETE_PRODUCT", "Products",
                                 f"Ẩn sản phẩm {sel[0]}")
            self._load()

    def _export_excel(self):
        if not OPENPYXL:
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile=f"SanPham_{datetime.now().strftime('%Y%m%d')}.xlsx",
        )
        if not path:
            return
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sản phẩm"
        headers = ["Mã SP","Tên","Danh mục","Thương hiệu","Giá nhập","Giá bán","Tồn kho","Bảo hành(tháng)"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="1a73e8")
            cell.alignment = Alignment(horizontal="center")
        rows = self.db.fetchall("""
            SELECT p.code, p.name, c.name as cat, p.brand,
                   p.import_price, p.selling_price, p.quantity, p.warranty_months
            FROM products p LEFT JOIN categories c ON p.category_id=c.id
            WHERE p.is_active=1 ORDER BY p.code""")
        for i, r in enumerate(rows, 2):
            ws.append([r["code"], r["name"], r["cat"] or "", r["brand"] or "",
                       r["import_price"], r["selling_price"],
                       r["quantity"], r["warranty_months"]])
            if i % 2 == 0:
                for col in range(1, len(headers) + 1):
                    ws.cell(row=i, column=col).fill = PatternFill(fill_type="solid", fgColor="e8f0fe")
        wb.save(path)
        messagebox.showinfo("Thành công", f"Đã xuất {len(rows)} sản phẩm ra file Excel!")
